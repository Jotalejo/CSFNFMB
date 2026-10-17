from datetime import date
from decimal import Decimal
from math import ceil

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import HTMLResponse
from sqlalchemy import text
from sqlalchemy.orm import Session

from dependencies import get_db, templates

from io import BytesIO
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


router = APIRouter(prefix="/reportes", tags=["reportes"])

MAX_FILAS = 5000
CLIENTES_POR_PAGINA = 100


SQL_SIN_RECOLECCION = text("""
    SELECT
        c.cod_cli AS codigo_cliente,
        c.razonsoc_cli AS cliente
    FROM clientes AS c
    WHERE NOT EXISTS (
        SELECT 1
        FROM recoleccion AS r
        WHERE r.codcli_recolec = c.cod_cli
          AND r.fecha_recolec BETWEEN :desde AND :hasta
    )
    ORDER BY c.razonsoc_cli, c.cod_cli
    LIMIT 5001
""")


SQL_RESIDUOS = text("""
    SELECT
        c.cod_cli AS codigo_cliente,
        c.razonsoc_cli AS cliente,
        t.cod_tipores AS codigo_residuo,
        t.nomtipo_tipores AS tipo_residuo,
        SUM(
            CAST(COALESCE(d.pesotot_detrec, 0) AS DECIMAL(20, 6))
        ) AS total
    FROM recoleccion AS r
    INNER JOIN clientes AS c
        ON c.cod_cli = r.codcli_recolec
    INNER JOIN detalle_recoleccion AS d
        ON d.codrecolec_detrec = r.cod_recolec
    INNER JOIN tiposresid AS t
        ON t.cod_tipores = d.codtipores_detrec
    WHERE r.fecha_recolec BETWEEN :desde AND :hasta
    GROUP BY
        c.cod_cli,
        c.razonsoc_cli,
        t.cod_tipores,
        t.nomtipo_tipores
    ORDER BY
        c.razonsoc_cli,
        c.cod_cli,
        t.nomtipo_tipores,
        t.cod_tipores
    LIMIT 5001
""")


def mostrar_reporte(
    request,
    db,
    tipo,
    fecha_desde,
    fecha_hasta,
    consultar,
    pagina,
    exportar=False,
):
    es_residuos = tipo == "residuos"
    titulo = (
        "Residuos por cliente"
        if es_residuos
        else "Clientes sin recolección"
    )

    hoy = date.today()
    contexto = {
        "request": request,
        "titulo": titulo,
        "es_residuos": es_residuos,
        "fecha_desde": (
            fecha_desde.isoformat()
            if fecha_desde
            else hoy.replace(day=1).isoformat()
        ),
        "fecha_hasta": (
            fecha_hasta.isoformat() if fecha_hasta else hoy.isoformat()
        ),
        "consultado": consultar,
        "error": "",
        "grupos": [],
        "total_clientes": 0,
        "pagina": 1,
        "paginas": 0,
        "anterior_url": None,
        "siguiente_url": None,
    }

    if consultar:
        if not fecha_desde or not fecha_hasta:
            contexto["error"] = "Debes seleccionar ambas fechas."
        elif fecha_desde > fecha_hasta:
            contexto["error"] = (
                "La fecha inicial no puede ser mayor que la final."
            )
        else:
            sql = SQL_RESIDUOS if es_residuos else SQL_SIN_RECOLECCION

            filas = db.execute(
                sql,
                {"desde": fecha_desde, "hasta": fecha_hasta},
            ).mappings().all()

            cantidad_clientes = len({
                fila["codigo_cliente"] for fila in filas
            })
            filas_reporte = len(filas) + (
                cantidad_clientes if es_residuos else 0
            )

            if filas_reporte > MAX_FILAS:
                contexto["error"] = (
                    "El resultado supera las 5.000 filas. "
                    "Reduce el rango de fechas y vuelve a consultar."
                )
            else:
                grupos = []
                por_cliente = {}

                for fila in filas:
                    codigo = fila["codigo_cliente"]

                    if codigo not in por_cliente:
                        grupo = {
                            "cliente": fila["cliente"],
                            "residuos": [],
                            "total": Decimal("0"),
                        }
                        por_cliente[codigo] = grupo
                        grupos.append(grupo)

                    if es_residuos:
                        total = fila["total"] or Decimal("0")
                        grupo = por_cliente[codigo]
                        grupo["residuos"].append({
                            "tipo": fila["tipo_residuo"],
                            "total": total,
                        })
                        grupo["total"] += total
                
                if exportar:
                    return exportar_excel(
                        grupos,
                        titulo,
                        es_residuos,
                        fecha_desde,
                        fecha_hasta,
                    )

                paginas = ceil(len(grupos) / CLIENTES_POR_PAGINA)
                pagina = min(pagina, max(paginas, 1))
                inicio = (pagina - 1) * CLIENTES_POR_PAGINA

                contexto.update({
                    "grupos": grupos[
                        inicio:inicio + CLIENTES_POR_PAGINA
                    ],
                    "total_clientes": len(grupos),
                    "pagina": pagina,
                    "paginas": paginas,
                })

                def url_pagina(numero):
                    return str(request.url.include_query_params(
                        fecha_desde=fecha_desde.isoformat(),
                        fecha_hasta=fecha_hasta.isoformat(),
                        consultar="true",
                        pagina=numero,
                    ))

                if pagina > 1:
                    contexto["anterior_url"] = url_pagina(pagina - 1)

                if pagina < paginas:
                    contexto["siguiente_url"] = url_pagina(pagina + 1)

    return templates.TemplateResponse(
        "reportes/reporte.html",
        contexto,
    )

def exportar_excel(grupos, titulo, es_residuos, desde, hasta):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"

    columnas = 3 if es_residuos else 1

    hoja.cell(1, 1, titulo)
    hoja.cell(
        2, 1,
        f"Período: {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"
    )

    if columnas > 1:
        hoja.merge_cells(
            start_row=1, start_column=1,
            end_row=1, end_column=columnas,
        )
        hoja.merge_cells(
            start_row=2, start_column=1,
            end_row=2, end_column=columnas,
        )

    hoja["A1"].font = Font(bold=True, size=14)

    encabezados = (
        ["Cliente", "Tipo de residuo", "Total kg"]
        if es_residuos
        else ["Cliente"]
    )

    for columna, nombre in enumerate(encabezados, start=1):
        celda = hoja.cell(4, columna, nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="245B45")

    # Guarda los nombres como texto, incluso si comienzan con "=".
    def escribir_texto(fila, columna, valor):
        celda = hoja.cell(fila, columna, str(valor or ""))
        celda.data_type = "s"
        celda.alignment = Alignment(
            vertical="top", wrap_text=True
        )

    fila = 5

    for grupo in grupos:
        if es_residuos:
            for indice, residuo in enumerate(grupo["residuos"]):
                escribir_texto(
                    fila, 1,
                    grupo["cliente"] if indice == 0 else "",
                )
                escribir_texto(fila, 2, residuo["tipo"])

                hoja.cell(fila, 3, float(residuo["total"]))
                hoja.cell(fila, 3).number_format = "#,##0.00"
                fila += 1

            hoja.cell(fila, 2, "TOTAL CLIENTE")
            hoja.cell(fila, 3, float(grupo["total"]))
            hoja.cell(fila, 3).number_format = "#,##0.00"

            for celda in hoja[fila]:
                celda.font = Font(bold=True)
                celda.fill = PatternFill(
                    "solid", fgColor="E2EFDA"
                )

            fila += 1
        else:
            escribir_texto(fila, 1, grupo["cliente"])
            fila += 1

    hoja.column_dimensions["A"].width = 55

    if es_residuos:
        hoja.column_dimensions["B"].width = 35
        hoja.column_dimensions["C"].width = 18

    hoja.freeze_panes = "A5"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.print_title_rows = "1:4"

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    nombre = (
        "residuos_por_cliente"
        if es_residuos
        else "clientes_sin_recoleccion"
    )
    nombre = f"{nombre}_{desde}_{hasta}.xlsx"

    return StreamingResponse(
        archivo,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{nombre}"'
        },
    )

@router.get("/sin-recoleccion", response_class=HTMLResponse)
def clientes_sin_recoleccion(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    consultar: bool = False,
    exportar: bool = False,
    pagina: int = Query(1, ge=1),
    db: Session = Depends(get_db),
):
    return mostrar_reporte(
        request, db, "sin-recoleccion",
        fecha_desde, fecha_hasta,
        consultar or exportar, pagina, exportar,
    )


@router.get("/residuos-por-cliente", response_class=HTMLResponse)
def residuos_por_cliente(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    consultar: bool = False,
    exportar: bool = False,
    pagina: int = Query(1, ge=1),
    db: Session = Depends(get_db),
):
    return mostrar_reporte(
        request, db, "residuos",
        fecha_desde, fecha_hasta,
        consultar or exportar, pagina, exportar,
    )